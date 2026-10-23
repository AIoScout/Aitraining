#!/usr/bin/env python3
"""Build the annotated Student Lab Manual for the Car License Project (3-Char).

What this script does
---------------------
1. Launches the real Streamlit app headlessly (streamlit run app.py).
2. Drives a real sign_ocr session with Playwright: Start Page -> A-1-2
   workspace -> duplicate-character demo -> Settings / EMNIST Auto-Fill ->
   capture drawer -> Sample Worktable -> real training -> Synth Match /
   Synth Mismatch -> Export / Save view.
3. Injects numbered red callout boxes (``#e11d48``) through
   ``getBoundingClientRect()`` and captures 1440x900 @2x screenshots.
   The Sample Worktable capture is polarity-guarded: the build fails if the
   Processed Preview is not black ink on pure white (255).
4. Assembles a STANDALONE, software-only student manual with python-docx:
   Title / A) Objective / Briefing, then sections 2-6 walking the
   TFLiteTraining app workflow (open project -> four areas -> target sign &
   samples -> Edit-page preprocessing -> training -> preview & export),
   following the bullet style and table formats of the reference manual's
   Task 2.  No hardware walkthrough, no Arduino IDE setup, no MCU
   flashing — the manual is ONLY the app workflow.
5. Writes the .docx to the repository root and to the OneDrive Student_Helper
   folder.  The reference "TF Lite Training Manual.docx" is never opened,
   read or modified.

Usage
-----
    cd AItraining
    .venv/bin/python scripts/build_car_sign_manual.py

    # Reuse already-captured screenshots and only rebuild the Word file:
    .venv/bin/python scripts/build_car_sign_manual.py --only-docx

Requires: streamlit, playwright (+ system Chrome), pillow, python-docx.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import time
import urllib.request
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from PIL import Image, ImageDraw, ImageFont
from playwright.sync_api import sync_playwright

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - actionable message for students
    raise SystemExit(
        "python-docx is required: .venv/bin/python -m pip install python-docx"
    ) from exc


APP_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = APP_DIR.parent
ASSET_DIR = APP_DIR / "docs" / "manual_assets" / "sign_ocr"
BUILD_LOG = ASSET_DIR / "build.log"
LOCAL_DOCX = REPO_ROOT / "Lab_Manual_Car_Sign.docx"
ONEDRIVE_DOCX = Path(
    "/Users/koil/Library/CloudStorage/OneDrive-HKUSTConnect/"
    "Student_Helper/TF Lite Training Manual - Car Sign.docx"
)

VIEWPORT = {"width": 1440, "height": 900}
DEVICE_SCALE = 2
CALL_OUT = "#e11d48"
BODY_FONT = "Arial"
HEADING_FONT = "Arial"


# ── logging / process helpers ───────────────────────────────────────────────
def log(message: str) -> None:
    stamp = time.strftime("%H:%M:%S")
    line = f"[{stamp}] {message}"
    print(line, flush=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with BUILD_LOG.open("a", encoding="utf-8") as fh:
        fh.write(line + "\n")


def free_port() -> int:
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    sock.close()
    return port


def wait_http(url: str, timeout_s: float = 90.0) -> None:
    deadline = time.time() + timeout_s
    last_error: Optional[Exception] = None
    while time.time() < deadline:
        try:
            with urllib.request.urlopen(url, timeout=2) as res:
                if res.status < 500:
                    return
        except Exception as exc:  # noqa: BLE001 - startup polling
            last_error = exc
        time.sleep(0.4)
    raise RuntimeError(f"server did not become ready at {url}: {last_error}")


class StreamlitApp:
    def __init__(self, port: Optional[int] = None) -> None:
        self.port = int(port or free_port())
        self.url = f"http://127.0.0.1:{self.port}/"
        self.proc: Optional[subprocess.Popen] = None
        self.log_file = None

    def start(self) -> None:
        ASSET_DIR.mkdir(parents=True, exist_ok=True)
        self.log_file = open(ASSET_DIR / "streamlit.log", "w", encoding="utf-8")
        env = dict(os.environ)
        env.setdefault("TF_CPP_MIN_LOG_LEVEL", "3")
        env.setdefault("STREAMLIT_BROWSER_GATHER_USAGE_STATS", "false")
        self.proc = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "streamlit",
                "run",
                "app.py",
                "--server.headless",
                "true",
                "--server.address",
                "127.0.0.1",
                "--server.port",
                str(self.port),
                "--browser.gatherUsageStats",
                "false",
                "--server.fileWatcherType",
                "none",
            ],
            cwd=str(APP_DIR),
            env=env,
            stdout=self.log_file,
            stderr=subprocess.STDOUT,
        )
        wait_http(f"{self.url}_stcore/health")
        log(f"Streamlit ready at {self.url}")

    def stop(self) -> None:
        if self.proc is not None:
            self.proc.terminate()
            try:
                self.proc.wait(timeout=12)
            except Exception:  # noqa: BLE001 - best-effort teardown
                self.proc.kill()
                try:
                    self.proc.wait(timeout=5)
                except Exception:
                    pass
        if self.log_file is not None:
            self.log_file.close()
        self.proc = None
        self.log_file = None


# ── browser helpers ─────────────────────────────────────────────────────────
WS_FRAME_JS = """
() => Array.from(document.querySelectorAll('iframe')).find((f) => {
  try { return !!(f.contentDocument && f.contentDocument.querySelector('#classes')); }
  catch (e) { return false; }
})
"""

MAIN_SCROLL_JS = """() => {
  const main = document.querySelector('section[data-testid="stMain"]');
  const rect = main.getBoundingClientRect();
  return { top: rect.top, scrollTop: main.scrollTop, height: main.clientHeight };
}"""

IFRAME_CONTENT_TOP_JS = """
() => {
  const main = document.querySelector('section[data-testid="stMain"]');
  const frame = Array.from(document.querySelectorAll('iframe')).find((f) => {
    try { return !!(f.contentDocument && f.contentDocument.querySelector('#classes')); }
    catch (e) { return false; }
  });
  if (!main || !frame) return null;
  return frame.getBoundingClientRect().top - main.getBoundingClientRect().top + main.scrollTop;
}"""


def wait_workspace(page) -> Any:
    page.wait_for_function(
        "() => Array.from(document.querySelectorAll('iframe')).some((f) => {"
        " try { return !!(f.contentDocument && f.contentDocument.querySelector('#classes')); }"
        " catch (e) { return false; } })",
        timeout=90_000,
    )
    frame = next(
        f
        for f in page.frames
        if f != page.main_frame and f.locator("#classes").count() > 0
    )
    frame.wait_for_selector("#signSeqBar", timeout=30_000)
    return frame


def dismiss_tour(ws) -> None:
    """Close the workspace tour before capturing.

    The tour is a checkpoint-driven walk: `#tmTourClose` (the × in the card
    head) is the exit and `#tmTourAutoskip` keeps later steps from walking
    the student through the UI while screenshots are being taken.  The tour
    may also start a moment after the workspace renders, so wait for it
    briefly instead of checking only once.
    """
    deadline = time.time() + 5.0
    while time.time() < deadline:
        try:
            if ws.locator("#tmTourClose").count():
                ws.evaluate(
                    """() => {
                      const auto = document.getElementById('tmTourAutoskip');
                      if (auto && !auto.checked) auto.click();
                      const close = document.getElementById('tmTourClose');
                      if (close) close.click();
                    }"""
                )
                log("Tour closed")
                time.sleep(0.4)
                return
        except Exception:  # noqa: BLE001 - tour may already be gone
            pass
        time.sleep(0.35)
    log("Tour not present (or already closed)")


def close_training_warning(ws) -> None:
    ws.evaluate(
        """() => {
          const nodes = Array.from(document.querySelectorAll('div'))
            .filter((d) => d.textContent.includes('Training Warnings') && d.querySelector('button'))
            .sort((a, b) => a.textContent.length - b.textContent.length);
          if (nodes.length) nodes[0].querySelector('button').click();
        }"""
    )
    time.sleep(0.4)


def scroll_main_content(page, content_top: float, target: float = 60.0) -> None:
    page.evaluate(
        """([contentTop, target]) => {
          const main = document.querySelector('section[data-testid="stMain"]');
          main.scrollTop = Math.max(0, contentTop - target);
        }""",
        [content_top, target],
    )
    time.sleep(0.25)


def scroll_to_iframe_y(page, iframe_y: float, target: float = 0.0) -> None:
    content_top = page.evaluate(IFRAME_CONTENT_TOP_JS)
    if content_top is None:
        raise RuntimeError("workspace iframe not found")
    scroll_main_content(page, float(content_top) + float(iframe_y), target)


def ws_top(ws, selector: str) -> float:
    return float(
        ws.evaluate(
            "(sel) => { const e = document.querySelector(sel);"
            " if (!e) throw new Error('missing ' + sel);"
            " return e.getBoundingClientRect().top; }",
            selector,
        )
    )


def ws_union(ws, selectors: Sequence[str]) -> Dict[str, float]:
    return ws.evaluate(
        """(sels) => {
          let left = 1e9, top = 1e9, right = -1e9, bottom = -1e9;
          for (const sel of sels) {
            const e = document.querySelector(sel);
            if (!e) throw new Error('missing callout target: ' + sel);
            const r = e.getBoundingClientRect();
            left = Math.min(left, r.left); top = Math.min(top, r.top);
            right = Math.max(right, r.right); bottom = Math.max(bottom, r.bottom);
          }
          return { left, top, right, bottom };
        }""",
        list(selectors),
    )


def add_callouts(frame, specs: Sequence[Dict[str, Any]]) -> None:
    """Draw numbered red boxes around one or more DOM selectors per number."""
    frame.evaluate(
        """(specs) => {
          document.querySelectorAll('.__callout_box, .__callout_badge').forEach((e) => e.remove());
          for (const spec of specs) {
            const elements = spec.selectors.map((s) => document.querySelector(s)).filter(Boolean);
            if (!elements.length) throw new Error('callout target missing: ' + spec.selectors.join(', '));
            let left = 1e9, top = 1e9, right = -1e9, bottom = -1e9;
            for (const el of elements) {
              const r = el.getBoundingClientRect();
              left = Math.min(left, r.left); top = Math.min(top, r.top);
              right = Math.max(right, r.right); bottom = Math.max(bottom, r.bottom);
            }
            const pad = spec.pad == null ? 6 : Number(spec.pad);
            const box = document.createElement('div');
            box.className = '__callout_box';
            Object.assign(box.style, {
              position: 'fixed',
              left: (left - pad) + 'px',
              top: (top - pad) + 'px',
              width: (right - left + 2 * pad) + 'px',
              height: (bottom - top + 2 * pad) + 'px',
              border: '3px solid #e11d48',
              borderRadius: '9px',
              boxSizing: 'border-box',
              zIndex: 2147483600,
              pointerEvents: 'none',
              boxShadow: '0 0 0 1px rgba(255,255,255,0.65) inset'
            });
            const badge = document.createElement('div');
            badge.className = '__callout_badge';
            badge.textContent = String(spec.number);
            const badgeInside = !!spec.badgeInside;
            Object.assign(badge.style, {
              position: 'fixed',
              left: (badgeInside ? left - pad + 4 : Math.max(2, left - pad - 13)) + 'px',
              top: (badgeInside ? top - pad + 4 : Math.max(2, top - pad - 13)) + 'px',
              width: '26px',
              height: '26px',
              borderRadius: '50%',
              background: '#e11d48',
              color: '#ffffff',
              font: '700 14px/26px Arial, sans-serif',
              textAlign: 'center',
              zIndex: 2147483601,
              pointerEvents: 'none',
              boxShadow: '0 0 0 2px #ffffff'
            });
            document.body.appendChild(box);
            document.body.appendChild(badge);
          }
        }""",
        list(specs),
    )
    time.sleep(0.15)


def clear_callouts(frame) -> None:
    frame.evaluate(
        "() => document.querySelectorAll('.__callout_box, .__callout_badge').forEach((e) => e.remove())"
    )


def snap(page, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(path), full_page=False)
    log(f"screenshot -> {path.name}")


def _verify_processed_preview_polarity(ws) -> None:
    """Guard the one screenshot whose polarity matters pedagogically.

    The manual explains that the Sample Worktable's Processed Preview shows
    the dark character on a pure white (255) background.  If the app ever
    renders the old inverted white-on-black glyph, fail the build instead of
    shipping a screenshot that contradicts the text.
    """
    shot = ws.locator(".class-preprocess-result img").screenshot()
    im = Image.open(io.BytesIO(shot)).convert("L")
    px = list(im.getdata())
    n = len(px) or 1
    mean = sum(px) / n
    dark = sum(1 for p in px if p < 100) / n
    bright = sum(1 for p in px if p > 200) / n
    log(
        f"processed-preview polarity: mean={mean:.0f} "
        f"dark={dark * 100:.1f}% bright={bright * 100:.1f}%"
    )
    if not (mean > 150 and 0.005 <= dark <= 0.60 and bright > 0.30):
        raise RuntimeError(
            "Sample Worktable Processed Preview is not black ink on white "
            "(polarity guard) — the 05 screenshot would mislead students"
        )


def click_js(ws, expression: str) -> None:
    ws.evaluate(expression)
    time.sleep(0.25)


# ── real-app walkthrough ────────────────────────────────────────────────────
def capture_walkthrough(app: StreamlitApp) -> Dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    shots: Dict[str, Path] = {}

    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel="chrome", headless=True)
        # Force prefers-color-scheme: light — the app follows the OS dark
        # mode setting, and a printed student manual wants the same clean
        # light UI as the reference manual's screenshots.
        page = browser.new_page(
            viewport=VIEWPORT,
            device_scale_factor=DEVICE_SCALE,
            color_scheme="light",
        )
        page.goto(app.url, wait_until="domcontentloaded", timeout=90_000)
        page.wait_for_selector(".tm-start-card", timeout=120_000)
        page.wait_for_timeout(1200)

        # ── 01 Start Page ──────────────────────────────────────────────────
        page.wait_for_timeout(500)
        card = page.locator(".tm-start-card").nth(2)
        get_started = page.get_by_role("button", name="Get Started").nth(2)
        open_project = page.get_by_role("button", name="Open Project (.tmproj)").nth(2)
        card.scroll_into_view_if_needed()
        page.wait_for_timeout(400)
        # Make sure the card and both buttons are in the 1440x900 viewport.
        open_box = open_project.bounding_box()
        if open_box and open_box["y"] + open_box["height"] > 850:
            page.evaluate(
                """(dy) => {
                  const main = document.querySelector('section[data-testid="stMain"]');
                  main.scrollTop += dy;
                }""",
                open_box["y"] + open_box["height"] - 830,
            )
            page.wait_for_timeout(350)
        boxes = {
            "card": card.bounding_box(),
            "start": get_started.bounding_box(),
            "open": open_project.bounding_box(),
        }
        # The Streamlit buttons have no stable ids; draw from measured boxes.
        page.evaluate(
            """(boxes) => {
              document.querySelectorAll('.__callout_box, .__callout_badge').forEach((e) => e.remove());
              const specs = [
                { n: 1, r: boxes.card },
                { n: 2, r: boxes.start },
                { n: 3, r: boxes.open },
              ];
              for (const { n, r } of specs) {
                if (!r) continue;
                const box = document.createElement('div');
                box.className = '__callout_box';
                Object.assign(box.style, {
                  position: 'fixed', left: (r.x - 5) + 'px', top: (r.y - 5) + 'px',
                  width: (r.width + 10) + 'px', height: (r.height + 10) + 'px',
                  border: '3px solid #e11d48', borderRadius: '9px',
                  boxSizing: 'border-box', zIndex: 2147483600, pointerEvents: 'none'
                });
                const badge = document.createElement('div');
                badge.className = '__callout_badge';
                badge.textContent = String(n);
                Object.assign(badge.style, {
                  position: 'fixed', left: Math.max(2, r.x - 32) + 'px',
                  top: Math.max(2, r.y - 16) + 'px', width: '26px', height: '26px',
                  borderRadius: '50%', background: '#e11d48', color: '#fff',
                  font: '700 14px/26px Arial, sans-serif', textAlign: 'center',
                  zIndex: 2147483601, pointerEvents: 'none',
                  boxShadow: '0 0 0 2px #fff'
                });
                document.body.appendChild(box);
                document.body.appendChild(badge);
              }
            }""",
            boxes,
        )
        page.wait_for_timeout(200)
        shots["01_start_page"] = ASSET_DIR / "01_start_page.png"
        snap(page, shots["01_start_page"])
        clear_callouts(page)

        # ── Enter the sign workspace ──────────────────────────────────────
        get_started.click()
        ws = wait_workspace(page)
        page.wait_for_timeout(1500)
        dismiss_tour(ws)

        def set_sequence(seq: List[str]) -> None:
            ws.evaluate("(s) => setSignSequence(s)", seq)
            page.wait_for_timeout(1800)

        set_sequence(["A", "1", "2"])

        # ── 02 Workspace overview ─────────────────────────────────────────
        scroll_to_iframe_y(page, 0, target=0)
        page.wait_for_timeout(300)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": ["#menuFileBtn", "#navSettings", "#navTutorial", "#navReturn"], "pad": 6},
                {"number": 2, "selectors": ["#signSeqBar"], "pad": 5},
                {"number": 3, "selectors": [".class-card"], "pad": 5},
                {"number": 4, "selectors": ["#trainCard"], "pad": 5},
                {"number": 5, "selectors": ["#previewCard"], "pad": 5},
            ],
        )
        shots["02_workspace_overview"] = ASSET_DIR / "02_workspace_overview.png"
        snap(page, shots["02_workspace_overview"])
        clear_callouts(ws)

        # ── 03 Target sequence + duplicate dedup ──────────────────────────
        set_sequence(["A", "1", "A"])
        scroll_to_iframe_y(page, 0, target=0)
        page.wait_for_timeout(300)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": ["#signSeqSlot-0", "#signSeqSlot-1", "#signSeqSlot-2"], "pad": 7},
                {"number": 2, "selectors": ["#signSeqLeft-1", "#signSeqRight-1"], "pad": 6},
                {"number": 3, "selectors": ["#signSeqBar .signseq-head .signseq-note"], "pad": 6},
                {"number": 4, "selectors": [".signslot-badge"], "pad": 6},
            ],
        )
        shots["03_target_sequence_and_dedup"] = ASSET_DIR / "03_target_sequence_and_dedup.png"
        snap(page, shots["03_target_sequence_and_dedup"])
        clear_callouts(ws)
        set_sequence(["A", "1", "2"])

        # ── 04a Settings: source + EMNIST auto-fill ───────────────────────
        click_js(ws, "openGlobalSettings()")
        ws.wait_for_timeout(500)
        click_js(ws, "document.getElementById('settingsAdvancedToggle').click()")
        ws.evaluate(
            """() => {
              const s = document.getElementById('settingsPreviewSource');
              s.value = 'upload';
              s.dispatchEvent(new Event('change', {bubbles: true}));
            }"""
        )
        click_js(
            ws,
            """() => {
              const c = document.getElementById('settingsAutoSeedEmnist');
              if (!c.checked) { c.checked = true; c.dispatchEvent(new Event('change', {bubbles: true})); }
            }""",
        )
        ws.wait_for_function(
            "STATE.counts.A >= 24 && STATE.counts['1'] >= 24 && STATE.counts['2'] >= 24",
            timeout=180_000,
        )
        page.wait_for_timeout(3200)  # let the auto-fill toast fade
        log(f"Auto-fill complete: {ws.evaluate('JSON.stringify(STATE.counts)')}")

        # Top view: Preview source selector.
        ws.evaluate("document.querySelector('#globalSettingsModal .tm-settings-card').scrollTop = 0")
        scroll_to_iframe_y(page, 0, target=0)
        page.wait_for_timeout(300)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": ["#settingsPreviewSource"], "pad": 6},
            ],
        )
        settings_top = ASSET_DIR / "_settings_top.png"
        snap(page, settings_top)
        clear_callouts(ws)

        # Bottom view: Advanced -> EMNIST samples + Save.  The modal card is
        # taller than the visible iframe slice, so scroll BOTH the card to its
        # maximum and the Streamlit page to the bottom; the Advanced group and
        # Save then sit in the visible slice.
        ws.evaluate(
            """() => {
              const card = document.querySelector('#globalSettingsModal .tm-settings-card');
              card.scrollTop = card.scrollHeight;
            }"""
        )
        page.evaluate(
            """() => {
              const main = document.querySelector('section[data-testid="stMain"]');
              main.scrollTop = main.scrollHeight;
            }"""
        )
        page.wait_for_timeout(400)
        add_callouts(
            ws,
            [
                {"number": 2, "selectors": [".tm-settings-check"], "pad": 4},
                {"number": 3, "selectors": ["#settingsSeedEmnistNowBtn"], "pad": 4},
                {"number": 4, "selectors": ["#settingsSave"], "pad": 2},
            ],
        )
        settings_bottom = ASSET_DIR / "_settings_bottom.png"
        snap(page, settings_bottom)
        clear_callouts(ws)

        click_js(ws, "document.getElementById('settingsSave').click()")
        page.wait_for_timeout(1200)
        if ws.evaluate("document.getElementById('globalSettingsModal').classList.contains('open')"):
            click_js(ws, "document.getElementById('settingsCancel').click()")
            page.wait_for_timeout(600)

        # ── 04b Capture drawer ────────────────────────────────────────────
        click_js(
            ws,
            """() => {
              const card = document.querySelector('.class-card');
              const tiles = card.querySelectorAll('.tm-action-toolbar .sample');
              tiles[2].click();  // Device
            }""",
        )
        page.wait_for_timeout(2200)
        card_top = ws_top(ws, ".class-card")
        scroll_to_iframe_y(page, card_top, target=35)
        page.wait_for_timeout(350)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": [".class-card .tm-action-toolbar"], "pad": 5},
                {"number": 2, "selectors": ["#sourceCapture-A"], "pad": 5},
                {"number": 3, "selectors": ["#sourceHold-A"], "pad": 5},
                {"number": 4, "selectors": ["#sourceClose-A"], "pad": 5},
                {"number": 5, "selectors": ["#samplesHost-A"], "pad": 3, "badgeInside": True},
            ],
        )
        shots["04b_capture_drawer"] = ASSET_DIR / "04b_capture_drawer.png"
        snap(page, shots["04b_capture_drawer"])
        clear_callouts(ws)
        click_js(ws, "document.getElementById('sourceClose-A').click()")
        page.wait_for_timeout(500)

        # ── 05 Sample Worktable ───────────────────────────────────────────
        click_js(ws, "document.querySelector('.class-card .class-preprocess-chip').click()")
        ws.wait_for_selector("#classPreprocessImage", timeout=20_000)
        ws.wait_for_selector(".class-preprocess-result img", timeout=30_000)
        page.wait_for_timeout(1200)
        scroll_to_iframe_y(page, 0, target=0)
        page.wait_for_timeout(300)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": ["#classPreprocessStage"], "pad": 5},
                {"number": 2, "selectors": [".class-preprocess-result"], "pad": 5},
                {"number": 3, "selectors": ["#classPreprocessMode"], "pad": 5},
                {"number": 4, "selectors": ["#classPreprocessDark", "#classPreprocessLum"], "pad": 6},
                {"number": 5, "selectors": ["#classPreprocessSave", "#classPreprocessCancel"], "pad": 5},
            ],
        )
        shots["05_sample_worktable_edit"] = ASSET_DIR / "05_sample_worktable_edit.png"
        snap(page, shots["05_sample_worktable_edit"])
        clear_callouts(ws)
        _verify_processed_preview_polarity(ws)
        click_js(ws, "document.getElementById('classPreprocessCancel').click()")
        page.wait_for_timeout(600)

        # ── 06 Train for real ─────────────────────────────────────────────
        click_js(ws, "document.getElementById('advBtn').click()")
        page.wait_for_timeout(500)
        click_js(ws, "document.getElementById('trainBtn').click()")
        log("Training started…")
        ws.wait_for_function(
            "() => { const e = document.getElementById('trainStatus');"
            " return e && e.textContent.includes('Model Trained'); }",
            timeout=1_200_000,
        )
        log(f"Training complete: {ws.locator('#trainStatus').inner_text()}")
        page.wait_for_timeout(1200)
        close_training_warning(ws)
        if ws.evaluate("document.getElementById('trainAdvPanel').style.display === 'none'"):
            click_js(ws, "document.getElementById('advBtn').click()")
            page.wait_for_timeout(400)

        train_top = ws_top(ws, "#trainCard")
        scroll_to_iframe_y(page, train_top, target=20)
        page.wait_for_timeout(350)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": ["#advInkPolarity"], "pad": 5},
                {"number": 2, "selectors": ["#advReadingOrder"], "pad": 5},
                {"number": 3, "selectors": ["#advUseRec"], "pad": 5},
                {"number": 5, "selectors": ["#trainBtn"], "pad": 5},
            ],
        )
        training_card = ASSET_DIR / "_training_card.png"
        snap(page, training_card)
        clear_callouts(ws)

        # OCR/sign hyperparameter panel lives in Settings → Training
        # Hyperparameters.  Capture it separately for the composite.
        click_js(ws, "openGlobalSettings()")
        page.wait_for_timeout(500)
        ws.evaluate(
            """() => {
              const card = document.querySelector('#globalSettingsModal .tm-settings-card');
              const target = document.getElementById('modalOcrSection');
              const cr = card.getBoundingClientRect();
              const tr = target.getBoundingClientRect();
              card.scrollTop += (tr.top - cr.top - 60);
            }"""
        )
        scroll_to_iframe_y(page, 0, target=0)
        page.wait_for_timeout(300)
        add_callouts(
            ws,
            [
                {"number": 4, "selectors": ["#advOcrPanel"], "pad": 6},
            ],
        )
        ocr_panel = ASSET_DIR / "_ocr_panel.png"
        snap(page, ocr_panel)
        clear_callouts(ws)
        click_js(ws, "document.getElementById('settingsCancel').click()")
        page.wait_for_timeout(600)

        # ── 07a / 07b Synth Match + Mismatch ──────────────────────────────
        if ws.evaluate("previewSource") != "upload":
            log("Preview source was not upload; switching through JS")
            ws.evaluate(
                """() => {
                  STATE.preview_source = 'upload';
                  previewSource = 'upload';
                  renderPreviewCard();
                }"""
            )
            page.wait_for_timeout(500)
        click_js(ws, "document.getElementById('previewSynthMatchBtn').click()")
        ws.wait_for_function(
            "() => document.querySelector('#previewOutput .signmatch')"
            " && document.querySelector('#previewOutput .signmatch').classList.contains('yes')",
            timeout=240_000,
        )
        page.wait_for_timeout(1200)
        preview_top = ws_top(ws, "#previewCard")
        scroll_to_iframe_y(page, preview_top, target=20)
        page.wait_for_timeout(350)
        preview_specs = [
            {"number": 1, "selectors": [".preview-synth-group"], "pad": 5},
            {"number": 2, "selectors": [".preview-controls"], "pad": 6},
            {"number": 3, "selectors": ["#previewImage"], "pad": 5},
            {"number": 4, "selectors": [".signmatch-main"], "pad": 5},
            {"number": 5, "selectors": [".signmatch-slots"], "pad": 5},
        ]
        add_callouts(ws, preview_specs)
        shots["07a_preview_synth_match"] = ASSET_DIR / "07a_preview_synth_match.png"
        snap(page, shots["07a_preview_synth_match"])
        clear_callouts(ws)

        click_js(ws, "document.getElementById('previewSynthMismatchBtn').click()")
        ws.wait_for_function(
            "() => document.querySelector('#previewOutput .signmatch')"
            " && document.querySelector('#previewOutput .signmatch').classList.contains('no')"
            " && document.querySelector('#previewOutput .signmatch-hint')",
            timeout=240_000,
        )
        page.wait_for_timeout(1200)
        add_callouts(ws, preview_specs)
        shots["07b_preview_synth_mismatch"] = ASSET_DIR / "07b_preview_synth_mismatch.png"
        snap(page, shots["07b_preview_synth_mismatch"])
        clear_callouts(ws)

        # ── 08 Export + Save ──────────────────────────────────────────────
        click_js(ws, "document.getElementById('menuFileBtn').click()")
        scroll_to_iframe_y(page, 0, target=0)
        page.wait_for_timeout(300)
        add_callouts(
            ws,
            [
                {"number": 1, "selectors": ["#exportBtn"], "pad": 5},
                {"number": 2, "selectors": ["#menuFileBtn", "#navSaveProject", "#navSaveAsProject"], "pad": 5},
            ],
        )
        shots["08_export_and_save"] = ASSET_DIR / "08_export_and_save.png"
        snap(page, shots["08_export_and_save"])
        clear_callouts(ws)

        browser.close()

    # Composite 04a (Preview source + Advanced EMNIST controls).
    shots["04a_settings_emnist_autofill"] = ASSET_DIR / "04a_settings_emnist_autofill.png"
    _pair_panels(
        [
            (settings_top, (300, 470, 760, 900),
             "Settings → Preview: choose Local Image File"),
            (settings_bottom, (300, 520, 1140, 900),
             "Settings → Advanced: Auto-fill 24 EMNIST samples"),
        ],
        shots["04a_settings_emnist_autofill"],
    )

    # Composite 06 (Training card + Settings → Training Hyperparameters).
    shots["06_training_and_advanced"] = ASSET_DIR / "06_training_and_advanced.png"
    _pair_panels(
        [
            (training_card, (640, 0, 980, 900), "Training card → Advanced"),
            (ocr_panel, (680, 0, 1140, 900), "Settings → Training Hyperparameters"),
        ],
        shots["06_training_and_advanced"],
    )

    # Remove intermediate captures; keep only the ten deliverable PNGs.
    for tmp in (settings_top, settings_bottom, training_card, ocr_panel):
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass
    return shots


# ── Pillow compositing ──────────────────────────────────────────────────────
def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    return ImageFont.load_default()


def _crop_css(image: Image.Image, css_box: Sequence[float]) -> Image.Image:
    scale = DEVICE_SCALE
    x1, y1, x2, y2 = [int(round(float(v) * scale)) for v in css_box]
    x1 = max(0, min(image.width - 1, x1))
    y1 = max(0, min(image.height - 1, y1))
    x2 = max(x1 + 1, min(image.width, x2))
    y2 = max(y1 + 1, min(image.height, y2))
    return image.crop((x1, y1, x2, y2))


def _pair_panels(
    panels: Sequence[tuple[Path, Sequence[float], str]],
    out_path: Path,
) -> None:
    """Two cropped panels side by side, each with a title strip.

    Used where one screenshot cannot show two distant parts of the UI
    (e.g. Settings → Preview vs Settings → Advanced Settings, or the
    Training card vs the Training Hyperparameters modal section).
    """
    canvas_w = VIEWPORT["width"] * DEVICE_SCALE
    canvas_h = VIEWPORT["height"] * DEVICE_SCALE
    canvas = Image.new("RGB", (canvas_w, canvas_h), "white")
    draw = ImageDraw.Draw(canvas)
    margin = 40
    gap = 50
    label_h = 62
    slot_w = (canvas_w - 2 * margin - gap) // 2
    slot_h = canvas_h - label_h - margin
    x_positions = [margin, margin + slot_w + gap]
    for idx, (path, crop_box, label) in enumerate(panels[:2]):
        crop = _crop_css(Image.open(path).convert("RGB"), crop_box)
        scale = min(slot_w / crop.width, slot_h / crop.height)
        crop = crop.resize((max(1, int(crop.width * scale)), max(1, int(crop.height * scale))),
                           Image.LANCZOS)
        x = x_positions[idx]
        draw.text((x, 14), label, font=_font(26, bold=True), fill=(15, 23, 42))
        canvas.paste(crop, (x, label_h))
    draw.line(
        [(margin + slot_w + gap // 2, 8), (margin + slot_w + gap // 2, canvas_h - 20)],
        fill=(203, 213, 225),
        width=2,
    )
    canvas.save(out_path)
    log(f"composite -> {out_path.name} ({canvas.width}x{canvas.height})")


# ── Word manual ─────────────────────────────────────────────────────────────
def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_borders(cell, color: str = "E11D48", size: int = 18) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.0,
                   color: Optional[str] = None, align: Optional[Any] = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = BODY_FONT
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]],
               widths: Optional[Sequence[float]] = None, font_size: float = 9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    _set_repeat_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        _set_cell_shading(cell, "E2E8F0")
        _set_cell_text(cell, text, bold=True, size=font_size, color="0F172A")
    for row_values in rows:
        row = table.add_row()
        # Never split a row across pages (the long "Use recommended" row
        # otherwise left an empty half-row at the top of the next page).
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
        for i, text in enumerate(row_values):
            _set_cell_text(row.cells[i], text, size=font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return table


def _callout_box(doc, title: str, body: Sequence[str], icon: str = "⚠") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "FEF2F2")
    _set_cell_borders(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(f"{icon} {title}")
    run.bold = True
    run.font.name = BODY_FONT
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
    for text in body:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = BODY_FONT
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x7F, 0x1D, 0x1D)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _figure(doc, path: Path, caption: str, width_in: float = 6.5) -> None:
    # Keep the lead-in paragraph, the image and its caption on one page where
    # possible; an orphaned step heading above a page break reads poorly.
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.keep_together = True
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    run.font.name = BODY_FONT


def _bullet(doc, text: str, bold_prefix: Optional[str] = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = BODY_FONT
        r.font.size = Pt(10.5)
        rest = text[len(bold_prefix):]
        r2 = p.add_run(rest)
        r2.font.name = BODY_FONT
        r2.font.size = Pt(10.5)
    else:
        r = p.add_run(text)
        r.font.name = BODY_FONT
        r.font.size = Pt(10.5)


def _numbered(doc, text: str) -> None:
    """Manually numbered list item.

    Word's built-in List Number style shares one numbering definition, so
    every list in the document continued from the previous list (Step 4A
    began at "4", Step 7 at "8").  Writing the number as text and resetting
    whenever the previous paragraph is not a numbered item keeps each list
    independent and predictable.
    """
    previous = doc.paragraphs[-1] if doc.paragraphs else None
    number = 1
    if previous is not None and previous.style is not None and previous.style.name == "Manual Number":
        match = re.match(r"^(\d+)\.\s", previous.text)
        if match:
            number = int(match.group(1)) + 1
    p = doc.add_paragraph(style="Manual Number")
    p.paragraph_format.space_after = Pt(2)
    # Keep a short numbered list together; the Step 7 five-item rule list
    # otherwise ended with a lone item on the next page.
    p.paragraph_format.keep_with_next = True
    p.add_run(f"{number}. ")
    r = p.add_run(text)
    r.font.name = BODY_FONT
    r.font.size = Pt(10.5)


def _para(doc, text: str, *, bold: bool = False, size: float = 10.5,
          color: Optional[str] = None, space_after: float = 6.0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = BODY_FONT
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def _heading(doc, text: str, level: int) -> None:
    p = doc.add_heading(level=level)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = HEADING_FONT
    if level == 2:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
    else:  # level 4 — numbered sub-steps inside our Task 2
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run.font.name = BODY_FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)



def build_docx(assets: Dict[str, Path]) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Inches(1))

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(4)
    for style_name in ("List Bullet", "List Number"):
        st = styles[style_name]
        st.font.name = BODY_FONT
        st.font.size = Pt(10.5)
    manual = styles.add_style("Manual Number", WD_STYLE_TYPE.PARAGRAPH)
    manual.base_style = styles["Normal"]
    manual.font.name = BODY_FONT
    manual.font.size = Pt(10.5)
    manual.paragraph_format.left_indent = Inches(0.3)
    manual.paragraph_format.first_line_indent = Inches(-0.22)
    manual.paragraph_format.space_after = Pt(2)

    # Header / footer.
    header = section.header.paragraphs[0]
    header.text = "TFLiteTraining App Manual · Car License Project (3-Char)"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.name = BODY_FONT
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Car License Project (3-Char) · ")
    fr.font.name = BODY_FONT
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    _add_page_number(footer)

    # ── Title block ─────────────────────────────────────────────────────
    _heading(doc, "TFLiteTraining App Manual: Car License Project (3-Char)", 2)
    _heading(doc, "A) Objective:", 3)
    _numbered(doc, "To get familiar with the TFLiteTraining desktop app.")
    _numbered(doc, "To collect character samples, preprocess images with ROI and Darkness/Luminosity thresholds, and train an AI model.")
    _numbered(doc, "To test 3-character road sign / car plate recognition (🟢 YES — MATCH / 🔴 NO — NOT MATCH) and export the trained model.")
    _para(doc, "Briefing:", bold=True)
    _para(
        doc,
        "In this task, you will learn how to train and export a TensorFlow Lite model through "
        "data collection, image preprocessing, model training, and performance preview. "
        "Specifically, you will train a model that can read a 3-character road sign / car sign "
        "in sequence (such as A 1 2) and tell you whether the sign in front of the camera "
        "matches your target (🟢 YES — MATCH or 🔴 NO — NOT MATCH). Please remember to save "
        "your project regularly (File ▾ → Save) to avoid losing your work.",
    )

    # ── 2. Opening the project & the four main areas ────────────────────
    _heading(doc, "2. Opening the Project & the Four Main Areas", 3)
    _para(
        doc,
        "Open the app. On the Start Page, find the 🛑 Car License Project (3-Char) card and "
        "click Get Started to start a new project, or click Open Project (.tmproj) to open a "
        "saved template/project file.",
    )
    _figure(doc, assets["01_start_page"], "Figure 1 — Start Page: choose the 🛑 Car License Project (3-Char) card.")
    _para(doc, "After opening the project, you will see four main areas:")
    _figure(doc, assets["02_workspace_overview"], "Figure 2 — The workspace: the four main areas.")
    _bullet(doc, "Target Sign Bar (Top Left) – for typing the 3 letters or numbers on your road sign in order (e.g., A, 1, 2).", bold_prefix="Target Sign Bar (Top Left)")
    _bullet(doc, "Sample Area (Left) – for collecting and preprocessing single-character photos for each character card.", bold_prefix="Sample Area (Left)")
    _bullet(doc, "Train Area (Middle) – for training the AI model by adjusting hyperparameters.", bold_prefix="Train Area (Middle)")
    _bullet(doc, "Preview Area (Right) – for testing the trained model on a 3-character sign and checking the 🟢 YES / 🔴 NO result.", bold_prefix="Preview Area (Right)")

    # ── 3. Sample Area Part A ───────────────────────────────────────────
    _heading(doc, "3. Sample Area — Part A: Setting the Target Sign & Collecting Samples", 3)
    _para(doc, "Setting the 3-Character Target Sign:", bold=True)
    _bullet(doc, "Click each of the three boxes in the Target Sign Bar and type your 3 characters (for example, A, 1, 2). Use the ◀ / ▶ buttons beneath a slot to swap neighbouring characters.")
    _bullet(doc, "If your sign uses the same character twice (like A 1 A), the app automatically keeps just one A card with a “Slots 1 & 3” badge so you only need to collect samples for A once!")
    _figure(doc, assets["03_target_sequence_and_dedup"], "Figure 3 — The Target Sign Bar: typing characters, the ◀ / ▶ swap buttons and the duplicate-character badge.")
    _para(doc, "Collecting Samples for Each Character Card:", bold=True)
    _callout_box(
        doc,
        "Key rule: one card = one character",
        [
            "Each character card collects photos of ONE single letter/number (not the whole "
            "3-letter sign at once). This way the AI learns each character separately, and "
            "during Preview the app slices the whole sign into 3 boxes and checks each one "
            "by itself.",
        ],
        icon="💡",
    )
    _bullet(doc, "Click “Webcam” — uses your laptop or USB camera. Capture single photos by clicking “Capture”, or take photos continuously by holding “Hold to Capture” (or holding the Space bar). Click “Close” to fold the panel.", bold_prefix="Click “Webcam”")
    _bullet(doc, "Click “Upload” — choose one or more image files from your computer.", bold_prefix="Click “Upload”")
    _bullet(doc, "Click “Device” — captures images streamed from the smart car camera over USB serial (select the serial port in ⚙ Settings).", bold_prefix="Click “Device”")
    _figure(doc, assets["04b_capture_drawer"], "Figure 4b — The capture drawer on a character card: Webcam / Upload / Device tiles, Capture / Hold to Capture and Close.")
    _bullet(doc, "Quick Auto-Fill Option (⚙ Settings): If you want to test the software right away without a camera, click ⚙ Settings in the top bar, expand ▸ Advanced Settings, and check Auto-fill empty classes with EMNIST template samples (24 / class) (or click Fill Empty Classes with EMNIST Samples Now), then click Save changes. Each empty character card immediately gets 24 handwriting samples!", bold_prefix="Quick Auto-Fill Option (⚙ Settings)")
    _figure(doc, assets["04a_settings_emnist_autofill"], "Figure 4a — ⚙ Settings: the Auto-fill 24 EMNIST samples option (left: Preview source, right: Advanced Settings).")

    # ── 4. Sample Area Part B: the Edit page ────────────────────────────
    _heading(doc, "4. Sample Area — Part B: Preprocessing the Samples (Edit Page)", 3)
    _para(
        doc,
        "On each character card header there are small icons: the Edit button opens the image "
        "processing page, and the eraser icon clears all the samples on that card.",
    )
    _para(doc, "The “Edit” page is composed of two parts: Image Viewer (Left) and Sample Worktable (Right):")
    _bullet(doc, "In the Image Viewer, we can see two images: Original Sample and Processed Preview.", bold_prefix="In the Image Viewer")
    _bullet(doc, "For the Original Sample: A green (Auto) or blue (Manual ROI) ROI (Region of Interest) box is shown. This ROI defines the area that will be cropped and passed for further processing.", bold_prefix="For the Original Sample")
    _bullet(doc, "For the Processed Preview: Pixels in the original image whose brightness falls below the Darkness Threshold (Dark Thr) or above the Luminosity Threshold (Lum Thr) are converted to pure white (Brightness: 255), leaving the dark character cleanly isolated on a white background.", bold_prefix="For the Processed Preview")
    _bullet(doc, "In the Sample Worktable, you can adjust the parameters to fine-tune the processing effect and minimize background noise:", bold_prefix="In the Sample Worktable")
    _figure(doc, assets["05_sample_worktable_edit"], "Figure 5 — The Edit page: Image Viewer (left) and Sample Worktable (right), with the dark letter on a pure white (255) Processed Preview.")
    _add_table(
        doc,
        ["Parameter", "Effects"],
        [
            ["Preprocess Mode", "1. Auto (find sign): automatically finds and crops the character in the image. 2. Manual ROI: crop the ROI manually by dragging your mouse on the original image."],
            ["Dark Thr (Darkness Threshold)", "Any pixel whose brightness is less than Dark Thr turns into pure white (255)."],
            ["Lum Thr (Luminosity Threshold)", "Any pixel whose brightness is greater than Lum Thr turns into pure white (255)."],
            ["ROI X1, Y1, X2, Y2 (Only for Manual ROI usage)", "The relative position (0.00 to 1.00) of the top-left (X1, Y1) and bottom-right (X2, Y2) corners of the crop box."],
        ],
        widths=[1.9, 4.6],
    )
    _para(doc, "Keyboard shortcuts:", bold=True)
    _bullet(doc, "Arrow keys — select the sample you want to process or check (move to the next/previous sample).")
    _bullet(doc, "Shift + Arrow keys — move the ROI box in Manual ROI mode.")
    _bullet(doc, "F1 — choose Auto preprocess mode. F2 — choose Manual ROI mode.")
    _bullet(doc, "D — delete the sample. Esc — exit the “Edit” view. S — save the change.")
    _para(doc, "Please remember to save after any changes!", bold=True, color="B91C1C", size=10)

    # ── 5. Train Area ───────────────────────────────────────────────────
    _heading(doc, "5. Train Area — Setting Hyperparameters & Training", 3)
    _para(
        doc,
        "Next, move to the Train Area. After image processing is complete, click ▸ Advanced "
        "to view or adjust the training hyperparameters. Below is the table for each "
        "hyperparameter:",
    )
    _add_table(
        doc,
        ["#", "Hyperparameter", "Default", "Range", "What it does"],
        [
            ["1", "Ink Polarity", "0 — Dark ink on bright bg", "0 / 1", "Which way round your letters are: normal dark-on-light writing (0), or bright letters on a dark sign (1)."],
            ["2", "Sign Reading Direction", "Left to Right", "L ➔ R / R ➔ L", "The order the app reads the three boxes of your sign."],
            ["3", "Image Size", "28", "8–512", "How many pixels wide the small letter picture is. 28×28 fits one letter perfectly — leave it."],
            ["4", "Batch Size", "32", "1 or more", "How many photos the AI practices with in one go."],
            ["5", "Epochs", "20", "1 or more", "How many times the AI reads through all your photos. For this project the app follows its own smart two-phase plan."],
            ["6", "Validation Split", "0.25", "0–0.5", "The part of your photos kept secret for testing."],
            ["7", "Learning Rate", "0.0016", "0 or more", "How big each learning step is. Small steps = careful but slow."],
            ["8", "Mix EMNIST into every class", "On", "on / off", "Mixes built-in handwriting examples with your photos so the AI learns from more variety."],
            ["9", "Background / noise class", "On", "on / off", "Adds a “no letter here” class, so the AI can tell you when a box is empty."],
            ["10", "Fine-tune pre-trained backbone", "On", "on / off", "Starts from a smart pre-learned letter reader so your model trains fast and accurately even with just 20 photos!"],
        ],
        widths=[0.3, 1.55, 1.3, 1.0, 2.35],
        font_size=8.5,
    )
    _para(doc, "What to do, step by step:", bold=True)
    _numbered(doc, "Click 💡 Use recommended to automatically fill in recommended settings for your dataset size.")
    _numbered(doc, "Click Train Embedded Model and wait for training to finish.")
    _numbered(doc, "Click Model Trained · tap for accuracy to view the per-class accuracy popup.")
    _numbered(doc, "If you change the target characters later, a ⚠ Stale badge reminds you to click Train Embedded Model again.")
    _figure(doc, assets["06_training_and_advanced"], "Figure 6 — The Train Area → Advanced (left) and ⚙ Settings → Training Hyperparameters (right).")

    # ── 6. Preview Area & export ────────────────────────────────────────
    _heading(doc, "6. Preview Area — Testing 3-Slot Sign Matching & Exporting", 3)
    _para(doc, "Select the Preview Image Input Source in ⚙ Settings → Preview:")
    _bullet(doc, "Device (MCU) — live pictures from the smart car camera.", bold_prefix="Device (MCU)")
    _bullet(doc, "Webcam — your laptop camera.", bold_prefix="Webcam")
    _bullet(doc, "Local Image File — for uploaded images and the 3 quick-test buttons 🧪 Synth Match, 🔀 Synth Reversed, ❌ Synth Mismatch.", bold_prefix="Local Image File")
    _para(doc, "The toggles and sliders on the Preview card:", bold=True)
    _bullet(doc, "Toggle Input ON to start live prediction.", bold_prefix="Toggle Input")
    _bullet(doc, "Toggle ROI: when unchecked, it shows the full image with the green ROI box drawn on it; when checked, it zooms in to show only the cropped image inside the ROI box.", bold_prefix="Toggle ROI")
    _bullet(doc, "Toggle Orig: switches between the threshold-cleaned black-on-white view and the original raw grayscale image.", bold_prefix="Toggle Orig")
    _bullet(doc, "Dark / Lum sliders: adjust brightness thresholds live so the sign letters stand out cleanly on a white background.", bold_prefix="Dark / Lum sliders")
    _para(doc, "Reading the 3-Slot Sign Match Output:", bold=True)
    _figure(doc, assets["07a_preview_synth_match"], "Figure 7a — 🧪 Synth Match: all three slots agree → 🟢 YES — MATCH.")
    _figure(doc, assets["07b_preview_synth_mismatch"], "Figure 7b — ❌ Synth Mismatch: one slot disagrees → 🔴 NO — NOT MATCH.")
    _bullet(doc, "The app slices the detected sign horizontally into 3 character slots in your chosen reading direction.")
    _bullet(doc, "Each slot card shows WANT <character>, the character read by the AI, the 28×28 image fed to the model, and the confidence percentage (%).")
    _bullet(doc, "When all 3 slots match your target characters with at least 60% confidence, the banner turns green: 🟢 YES — MATCH. If any slot reads a different character, is below 60%, or is blank (_background), the banner turns red: 🔴 NO — NOT MATCH and explains which slot failed.")
    _para(doc, "Exporting the Model & Saving Your Project:", bold=True)
    _bullet(doc, "Click Export Model at the top-right of the Preview card to export the C/C++ model files (tm_model_data.cpp/.h, model_resolver.h, model_settings.cpp/.h, tm.tflite, labels.txt).")
    _bullet(doc, "Click File ▾ → Save (or Save as) at the top-left menu to save your project as a .tmproj file.")
    _figure(doc, assets["08_export_and_save"], "Figure 8 — Export Model on the Preview card and File ▾ → Save / Save as.")
    _para(
        doc,
        "That's it — you collected letters, trained an AI to read them, tested a whole sign, "
        "and exported your model. Well done!",
        bold=True,
        color="1D4ED8",
    )

    # Scope guard: this is the standalone software-only Car Sign manual.
    # No hardware walkthrough (Task 0), no Arduino IDE setup (Task 1), no
    # MCU flashing (Task 3) — a future edit that slips them back in fails
    # the build instead of silently producing the wrong manual.
    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += "\n" + cell.text
    for banned in (
        "Task 0",
        "Task 1",
        "Task 3",
        "Arduino",
        "esp32_mannual",
        "Serial Monitor",
        "Signpost",
        "Battery Cabin",
        "Boot Button",
        "TF Card",
    ):
        if banned.lower() in all_text.lower():
            raise RuntimeError(f"manual scope violation: contains {banned!r}")
    if len(doc.inline_shapes) != 10:
        raise RuntimeError(
            f"expected exactly 10 figures, found {len(doc.inline_shapes)}"
        )

    LOCAL_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(LOCAL_DOCX))
    log(f"docx -> {LOCAL_DOCX}")
    ONEDRIVE_DOCX.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(LOCAL_DOCX, ONEDRIVE_DOCX)
    log(f"docx -> {ONEDRIVE_DOCX}")
    return LOCAL_DOCX

# ── main ────────────────────────────────────────────────────────────────────
def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only-docx", action="store_true",
                        help="skip Playwright and reuse existing screenshots")
    parser.add_argument("--port", type=int, default=None,
                        help="use a fixed Streamlit port instead of a free one")
    args = parser.parse_args(list(argv) if argv is not None else None)

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    if BUILD_LOG.exists() and not args.only_docx:
        BUILD_LOG.unlink()

    expected = [
        "01_start_page", "02_workspace_overview", "03_target_sequence_and_dedup",
        "04a_settings_emnist_autofill", "04b_capture_drawer", "05_sample_worktable_edit",
        "06_training_and_advanced", "07a_preview_synth_match",
        "07b_preview_synth_mismatch", "08_export_and_save",
    ]
    assets = {name: ASSET_DIR / f"{name}.png" for name in expected}

    if not args.only_docx:
        app = StreamlitApp(port=args.port)
        try:
            app.start()
            captured = capture_walkthrough(app)
            assets.update(captured)
        finally:
            app.stop()
    else:
        missing = [str(p) for p in assets.values() if not p.exists()]
        if missing:
            raise SystemExit("missing screenshots:\n  " + "\n  ".join(missing))

    missing = [str(p) for p in assets.values() if not p.exists()]
    if missing:
        raise SystemExit("manual build incomplete; missing:\n  " + "\n  ".join(missing))

    build_docx(assets)
    manifest = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "project_type": "sign_ocr",
        "screenshots": [str(p.relative_to(APP_DIR)) for p in assets.values()],
        "docx": [str(LOCAL_DOCX), str(ONEDRIVE_DOCX)],
    }
    (ASSET_DIR / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    log("Manual build complete")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
